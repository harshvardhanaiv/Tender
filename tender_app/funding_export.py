"""A Word file from the plain text of a Funding AI result (what the user sees and may have edited).

Text format, one item per line: "# " a title, "## " a section heading, "- " a bullet, anything else a paragraph;
blank lines are skipped."""
from __future__ import annotations

import io
import re

MAX_CHARS = 60000
_BAD_NAME = re.compile(r"[^A-Za-z0-9 _.-]+")


def file_name(title: str, suffix: str = ".docx") -> str:
    stem = _BAD_NAME.sub("", title or "").strip().replace(" ", "_").lstrip("._-")[:80] or "grant"
    return stem + suffix


def build_docx(title: str, text: str) -> bytes:
    from docx import Document

    doc = Document()
    if title:
        doc.add_heading(title[:200], level=1)
    for raw in (text or "")[:MAX_CHARS].splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            doc.add_paragraph(line.strip())
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
