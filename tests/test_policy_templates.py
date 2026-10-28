"""Run: python tests/test_policy_templates.py   (plain asserts; the API test needs the db container).

Policy templates (anti-bribery, modern slavery ...) written into the company's Documents with a review date, and the
Answer Bank starter packs: the prompts, the cleaning of the AI's answer, the Word file, credits and refunds, saved files."""
import io
import os
import sys
import time
import uuid
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
os.environ["ENABLE_EMAIL_SCHEDULER"] = "0"
os.environ["ENABLE_SCHEDULERS"] = "0"

from tender_app import answer_starters as ans  # noqa: E402
from tender_app import policy_templates as pt  # noqa: E402

RUN = uuid.uuid4().hex[:8]
USER = f"zz_policy_{RUN}"
NOPROFILE = f"zz_policy_{RUN}_np"


def test_library_covers_the_usual_documents():
    keys = set(pt.TEMPLATES)
    assert {"anti_bribery", "modern_slavery", "equality_diversity", "health_safety", "data_protection", "information_security",
            "environmental", "carbon_reduction_plan", "safeguarding", "business_continuity", "complaints", "social_value"} <= keys
    for t in pt.TEMPLATES.values():
        assert t["title"] and t["summary"] and len(t["outline"]) >= 5 and all(h and g for h, g in t["outline"]), t["key"]
        assert t["facts"][0]["key"] == "responsible" and len({f["key"] for f in t["facts"]}) == len(t["facts"]), t["key"]
    assert all("outline" not in x for x in pt.listing()) and len(pt.listing()) == len(pt.TEMPLATES)


def test_facts_and_review_period_are_cleaned():
    t = pt.TEMPLATES["health_safety"]
    assert pt.clean_facts(t, {"employees": "  24 \n staff ", "responsible": "", "evil": "x", "ico": "1"}) == {"employees": "24 staff"}
    assert pt.clean_facts(t, "junk") == {} and len(pt.clean_facts(t, {"employees": "x" * 900})["employees"]) == 400
    assert pt.review_months(24) == 24 and pt.review_months("6") == 6 and pt.review_months(7) == 12 and pt.review_months(None) == 12


def test_review_date_arithmetic():
    assert pt.add_months(date(2026, 10, 28), 12) == date(2027, 10, 28)
    assert pt.add_months(date(2026, 8, 31), 6) == date(2027, 2, 28), "a short month ends on its last day"
    assert pt.add_months(date(2028, 2, 29), 12) == date(2029, 2, 28)
    assert pt.add_months(date(2026, 12, 15), 6) == date(2027, 6, 15)


def test_prompt_carries_company_facts_and_forbids_inventing():
    t = pt.TEMPLATES["anti_bribery"]
    system, user = pt.build_prompt(t, "Acme Ltd", "We service boilers.", {"responsible": "Jane Smith, MD"})
    assert "Never invent" in system and "[add:" in system
    assert "Acme Ltd" in user and "We service boilers." in user and "Jane Smith, MD" in user and "Bribery Act 2010" in user
    assert "Gifts and hospitality" in user and "None given." not in user
    assert "None given." in pt.build_prompt(t, "", "", {})[1]


def test_answer_is_cut_down_to_sections():
    t = pt.TEMPLATES["anti_bribery"]
    c = pt.clean({"sections": [{"heading": "Purpose", "paragraphs": ["p1", 5, ""], "bullets": ["b"]}, {"heading": "", "paragraphs": ["x"]},
                               {"heading": "Empty", "paragraphs": []}, "junk"], "gaps": ["name", 3]}, t)
    assert c == {"sections": [{"heading": "Purpose", "paragraphs": ["p1"], "bullets": ["b"]}], "gaps": ["name"]}
    assert pt.clean("nope", t) == {"sections": [], "gaps": []} and pt.clean({"sections": "x"}, t)["sections"] == []


def test_word_file_has_the_control_block_and_the_sections():
    from docx import Document
    t = pt.TEMPLATES["anti_bribery"]
    content = {"sections": [{"heading": "Purpose and scope", "paragraphs": ["This applies to everyone."], "bullets": ["Employees"]}], "gaps": []}
    data = pt.build_docx(t, "Acme Ltd", content, date(2026, 10, 28), date(2027, 10, 28), "")
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    cells = [c.text for row in doc.tables[0].rows for c in row.cells]
    assert "Acme Ltd" in text and "Anti-bribery and corruption policy" in text and "Purpose and scope" in text and "This applies to everyone." in text
    assert "not legal advice" in text and "Employees" in text
    assert "28 October 2026" in cells and "28 October 2027" in cells and "[add: name and job title]" in cells and "1.0" in cells
    assert pt.file_name("Acme Ltd.", "Anti-bribery and corruption policy") == "Acme_Ltd_Anti_bribery_and_corruption_policy_v1.docx"
    assert pt.file_name("", "") == "Policy_v1.docx" or pt.file_name("", "").endswith("_v1.docx")


def test_starter_packs_are_honest_placeholders():
    assert {"general", "construction_fm", "it_digital", "professional_services", "care_support"} <= set(ans.PACKS)
    seen = set()
    for key, pack in ans.PACKS.items():
        assert pack["title"] and pack["summary"] and pack["items"], key
        for it in pack["items"]:
            assert it["category"] and it["question"].endswith(("?", ".")) and it["answer"], (key, it["question"])
            assert "[add" in it["answer"], f"{key}: an answer states nothing the company has not said: {it['question']}"
            assert it["answer"].count("[") == it["answer"].count("]"), (key, it["question"])
            assert it["question"].lower() not in seen, f"duplicate question {it['question']}"
            seen.add(it["question"].lower())
    assert len(ans.PACKS["general"]["items"]) >= 12 and ans.listing()[0]["count"] == len(ans.PACKS["general"]["items"])


def test_certification_guides_are_complete_and_link_to_real_templates():
    from tender_app import certifications as certs
    assert {"iso9001", "iso14001", "iso45001", "iso27001", "cyber_essentials", "ssip", "insurance"} <= set(certs.GUIDES)
    for key, g in certs.GUIDES.items():
        assert g["title"] and g["what"] and g["why"] and g["issuer"] and g["time"] and g["cost"] and g["renewal"], key
        assert len(g["steps"]) >= 5 and all(s["title"] and s["detail"] for s in g["steps"]) and g["needs"] and g["pitfalls"] and g["doc_words"], key
        assert all(k in pt.TEMPLATES for k in g["policies"]), f"{key} points at a policy that does not exist"
        assert all(l["url"].startswith("https://") for l in g["links"]), key
    assert "UKAS" in certs.GUIDES["iso9001"]["issuer"] and "12 months" in certs.GUIDES["cyber_essentials"]["renewal"]
    assert "indicative" in certs.NOTE


def test_documents_that_look_like_a_certificate_are_found():
    from tender_app import certifications as certs
    got = certs.match_documents(["ACME_QMS_ISO9001.pdf", "Cyber Essentials Plus 2026.pdf", "CHAS-accreditation.pdf", "Public Liability Insurance.pdf", "holiday.jpg", "ISO 27001 certificate.pdf"])
    assert got["iso9001"] == ["ACME_QMS_ISO9001.pdf"] and got["cyber_essentials"] == ["Cyber Essentials Plus 2026.pdf"] and got["ssip"] == ["CHAS-accreditation.pdf"]
    assert got["insurance"] == ["Public Liability Insurance.pdf"] and got["iso27001"] == ["ISO 27001 certificate.pdf"] and "iso14001" not in got
    assert next(g for g in certs.listing(got) if g["key"] == "iso9001")["have"] == ["ACME_QMS_ISO9001.pdf"] and certs.match_documents([]) == {}


# ─────────────────────────────────────── API ──────────────────────────────────────────────────────

import unittest


class Skip(unittest.SkipTest):
    pass


def test_api_generates_saves_and_refunds():
    try:
        from tender_app.db import get_db_connection
        conn = get_db_connection()
    except Exception as exc:  # noqa: BLE001
        raise Skip(f"no database: {exc}")
    import server
    from tender_app.credits import get_balance, grant_credits

    def wallet(user):
        c = get_db_connection()
        try:
            return get_balance(c, user)
        finally:
            c.close()

    cur = conn.cursor()
    cur.execute("INSERT INTO company_profiles (username, name, profile_text, is_default) VALUES (%s, 'ZZ Heating Ltd', 'We service gas boilers.', TRUE) RETURNING id", (USER,))
    profile_id = cur.fetchone()[0]
    conn.commit()
    original = server.call_chat_json
    seen = []

    def fake(system, user, **kw):
        seen.append(user)
        return {"sections": [{"heading": "Purpose and scope", "paragraphs": ["We do not tolerate bribery."], "bullets": []}], "gaps": ["name of the person responsible"]}

    server.call_chat_json = fake
    path = None
    try:
        app = server.create_app()

        def client(user):
            c = app.test_client()
            with c.session_transaction() as s:
                s.update(logged_in=True, username=user, email=f"{user}@test.invalid", last_activity=time.time(), csrf_token="t")
            return c

        def post(c, url, body):
            r = c.post(url, json=body, headers={"X-CSRF-Token": "t"})
            return r.status_code, (r.get_json(silent=True) or {})

        c = client(USER)
        g = get_db_connection()
        grant_credits(g, USER, 12, "test")
        g.close()
        st, r = c.get("/api/policy-templates").status_code, c.get("/api/policy-templates").get_json()
        assert st == 200 and r["cost"] == 3 and len(r["templates"]) == len(pt.TEMPLATES) and r["default_review_months"] == 12
        assert client(NOPROFILE).get("/api/policy-templates").status_code == 200
        base = "/api/policy-templates/anti_bribery/generate"
        assert post(c, "/api/policy-templates/nope/generate", {})[0] == 404
        st, r = post(client(NOPROFILE), base, {})
        assert st == 400 and "company profile" in r["error"] and wallet(NOPROFILE) == 0, "no profile: refused before any charge"

        st, r = post(c, base, {"profile_id": profile_id, "facts": {"responsible": "Jane Smith, MD", "evil": "x"}, "approved_by": "Jane Smith, Managing Director", "review_months": 24})
        assert st == 200 and r["cost"] == 3 and wallet(USER) == 9, (st, r)
        assert "Jane Smith, MD" in seen[0] and "ZZ Heating Ltd" in seen[0] and "evil" not in seen[0]
        d = r["document"]
        assert d["filename"] == "ZZ_Heating_Ltd_Anti_bribery_and_corruption_policy_v1.docx" and d["company_name"] == "ZZ Heating Ltd"
        assert d["issue_date"] == date.today().isoformat() and d["expiry_date"] == pt.add_months(date.today(), 24).isoformat()
        assert r["content"]["gaps"] == ["name of the person responsible"]
        cur.execute("SELECT filename, file_path, file_size, issue_date, expiry_date, username FROM company_profile_documents WHERE id = %s", (d["id"],))
        row = cur.fetchone()
        path = row[1]
        assert row[0] == d["filename"] and row[3] == d["issue_date"] and row[4] == d["expiry_date"] and row[5] == USER and os.path.getsize(path) == row[2] > 1000
        from docx import Document
        assert "Jane Smith, Managing Director" in [x.text for x in Document(path).tables[0].columns[1].cells]
        dl = c.get(d["download_url"])
        assert dl.status_code == 200 and len(dl.data) == row[2], "the saved file downloads from Documents"
        assert client(NOPROFILE).get(d["download_url"]).status_code in (403, 404), "another user cannot download it"
        listed = [x for x in c.get("/api/user-documents").get_json()["documents"] if x["id"] == d["id"] and x["type"] == "profile"]
        assert listed and listed[0]["expiry_date"] == d["expiry_date"], "it appears in Documents with its review date"

        def down(system, user, **kw):
            raise server.AIServiceError("The AI provider is unavailable", status_code=503)

        server.call_chat_json = down
        before = wallet(USER)
        st, r = post(c, base, {"profile_id": profile_id})
        assert st == 503 and "not been charged" in r["error"] and wallet(USER) == before, "a failed run is refunded"
        server.call_chat_json = lambda system, user, **kw: {"sections": []}
        st, r = post(c, base, {"profile_id": profile_id})
        assert st == 502 and wallet(USER) == before, "an empty answer is refunded and nothing is saved"
        cur.execute("SELECT COUNT(*) FROM company_profile_documents WHERE username = %s", (USER,))
        assert cur.fetchone()[0] == 1
        server.call_chat_json = fake
        g = get_db_connection()
        gc = g.cursor()
        gc.execute("UPDATE credit_wallet SET balance = 0 WHERE username = %s", (USER,))
        g.commit()
        g.close()
        st, r = post(c, base, {"profile_id": profile_id})
        assert st == 402 and r["code"] == "INSUFFICIENT_CREDITS", (st, r)

        cur.execute("INSERT INTO company_profile_documents (profile_id, username, filename, file_path, file_size) VALUES (%s, %s, 'ACME_ISO9001_certificate.pdf', '/x', 1)", (profile_id, USER))
        conn.commit()
        guides = c.get("/api/certifications").get_json()
        assert guides["ok"] and len(guides["guides"]) == len(__import__("tender_app.certifications", fromlist=["x"]).GUIDES) and "indicative" in guides["note"]
        assert next(g for g in guides["guides"] if g["key"] == "iso9001")["have"] == ["ACME_ISO9001_certificate.pdf"]
        assert next(g for g in client(NOPROFILE).get("/api/certifications").get_json()["guides"] if g["key"] == "iso9001")["have"] == [], "another user's files are not matched"
        assert wallet(USER) == 0, "the guides are free"

        # Answer Bank starter packs: free, added once, for the chosen company
        assert c.get("/api/answer-starters").get_json()["packs"][0]["key"] == "general"
        n = len(ans.PACKS["general"]["items"])
        st, r = post(c, "/api/answer-starters/general/add", {"profile_id": profile_id})
        assert st == 200 and r == {"ok": True, "added": n, "skipped": 0} and wallet(USER) == 0
        assert post(c, "/api/answer-starters/general/add", {"profile_id": profile_id})[1] == {"ok": True, "added": 0, "skipped": n}, "a pack is added once"
        st, r = post(c, "/api/answer-starters/care_support/add", {"profile_id": profile_id})
        assert r["added"] == len(ans.PACKS["care_support"]["items"])
        cur.execute("SELECT COUNT(*) FROM answer_bank WHERE username = %s AND profile_id = %s", (USER, profile_id))
        assert cur.fetchone()[0] == n + len(ans.PACKS["care_support"]["items"])
        assert post(c, "/api/answer-starters/nope/add", {})[0] == 404
        assert post(client(NOPROFILE), "/api/answer-starters/general/add", {"profile_id": profile_id})[0] == 400, "not someone else's company"
        st, r = post(client(NOPROFILE), "/api/answer-starters/general/add", {})
        assert st == 200 and r["added"] == n, "with no company the answers are general ones"
    finally:
        server.call_chat_json = original
        g = get_db_connection()
        gc = g.cursor()
        for table in ("credit_ledger", "credit_wallet", "company_profile_documents", "answer_bank", "company_profiles"):
            gc.execute(f"DELETE FROM {table} WHERE username = ANY(%s)", ([USER, NOPROFILE],))
        g.commit()
        g.close()
        conn.close()
        if path and os.path.exists(path):
            os.remove(path)


if __name__ == "__main__":
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    ran = 0
    for t in tests:
        try:
            t()
        except Skip as exc:
            print("skip", t.__name__, exc)
            continue
        ran += 1
        print("ok", t.__name__)
    print(f"{ran} passed")
